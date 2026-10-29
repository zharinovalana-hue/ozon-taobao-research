"""Refresh inputs and saved comparisons in the existing authored workbook."""
import json,os,tempfile
from pathlib import Path
from workbook_io import FormulaBook

def refresh(root,inputs,cases,results,output):
    root=Path(root);source=root/'models/Ozon_модель_кейса.xlsx'
    if not source.exists():raise FileNotFoundError('Download the existing model from models/ before refreshing')
    mapping=json.loads((root/'data/workbook_map.json').read_text(encoding='utf-8'))
    if len(cases)!=mapping['scenarioCount'] or [c['id'] for c in cases]!=[r['id'] for r in mapping['caseRows']]:
        raise ValueError('Scenario schema changed; workbook layout must be explicitly upgraded, not silently truncated')
    b=FormulaBook(source)
    for address,label in [('X5','Справочно: вес C S1'),('Y5','Справочно: вес F S1'),('Z5','Справочно: вес оба S1'),('AA5','Справочно: вес вне S1')]:b.set('Сценарии',address,label)
    for k,c in enumerate(cases[:3]):
        letter=chr(67+k)
        for j,x in enumerate(c['shares']):b.set('Параметры',letter+str(15+j),x)
        for r,key in [(19,'whale_cr'),(20,'global_cr_input'),(21,'owners_small')]:b.set('Параметры',letter+str(r),c[key])
    for i,c in enumerate(cases):
        row=mapping['caseRows'][i]['row'];weights=c.get('weights',inputs['weights'])[0]
        values=[i+1,c['name'],c['kind'],c['dau'],c['mau'],*c['shares'],c['global_cr_input'],c['whale_cr'],c['anchor'],
                c['owners_small'],c['overlap'],c['reach_scale'],c['behavior_penalty'],c['product_penalty'],c['loss_c'],c['loss_f'],
                c['income_delta'],c['revenue_per_buyer'],int(c['enforce_ozon']),c['threshold'],*weights]
        for j,value in enumerate(values,1):
            if i<3 and j in [6,7,8,9,10,11,13] or j>=24:continue  # linked to original inputs / cohort weight matrix
            from openpyxl.utils.cell import get_column_letter
            b.set('Сценарии',get_column_letter(j)+str(row),value)
        for k,ws in enumerate(c.get('weights',inputs['weights'])):
            for col,value in zip(['N','O','P','Q'],ws):b.set('Входы H',col+str(mapping['weightsFirstRow']+4*i+k),value)
    b.set('Параметры','B3',mapping['activeScenario'])
    b.set('Параметры','B10',inputs['uplift'])
    for h,definition in enumerate(inputs['hypotheses']):
        for k in range(4):
            r=5+4*h+k
            for col,value in zip(['C','D','E'],definition['reach'][k]):b.set('Входы H',col+str(r),value)
    for k in range(4):
        for j,col in enumerate(['H','I']):b.set('Когорты',col+str(k+5),inputs['traffic'][k][j])
    for addr,key in [('B5','case_daily_cart'),('B6','case_daily_fav'),('B7','case_monthly_cart'),('B8','case_monthly_fav')]:b.set('Исходные',addr,inputs[key])
    for addr,gap in zip(['B17','B18','B19'],inputs['case_cr_gaps'][1:]):b.set('Исходные',addr,gap*100)
    p=inputs['public_ozon']
    for r,key in enumerate(['dau_march_2026','mau_march_2026','q1_orders','q1_days','q1_active_buyers','trailing_orders_per_buyer','q3_orders_rounded','q3_frequency'],6):b.set('Масштаб','B'+str(r),p[key])
    for i,result in enumerate(results):
        r=mapping['comparisonFirstRow']+i
        for j,key in enumerate(mapping['comparisonFields'],1):
            value=result[key]
            if isinstance(value,bool):value='Да' if value else 'Нет'
            b.set('Сравнение',get_column_letter(j)+str(r),value)
    for h,definition in enumerate(inputs['hypotheses']):
        for j,id in enumerate(mapping['topCases'],3):
            x=next(r for r in results if r['scenario']==id and r['hypothesis']==definition['id'])
            b.set('Сравнение',get_column_letter(j)+str(7+h),x['delta_pp'])
    summary=json.loads((root/'data/taobao_summary.json').read_text(encoding='utf-8'))
    for i,key in enumerate(['users','cart_users','fav_users','both_users','buyers'],6):b.set('Taobao','B'+str(i),summary[key])
    r=14
    for t in summary['thresholds']:
        for x in t['cohorts']:
            vals=['P'+str(round(t['quantile']*100)),t['cart_threshold'],t['fav_threshold'],x['id'],x['users'],x['share'],x['buyers'],x['purchase_rate']]
            for j,val in enumerate(vals,1):b.set('Taobao',get_column_letter(j)+str(r),val)
            r+=1
    for capture in mapping['capture']:
        value=b.get(capture['sheet'],capture['address'])
        b.set('Версия срезов','C'+str(capture['row']),value)
    if b.get('Сравнение','B3')!='Актуально':
        mismatches=[(x['sheet'],x['address'],b.get('Версия срезов','C'+str(x['row'])),b.get('Версия срезов','D'+str(x['row']))) for x in mapping['capture'] if b.get('Версия срезов','E'+str(x['row']))!=1]
        raise ValueError('Snapshot capture did not reconcile: '+str(mismatches[:5]))
    for h,x in enumerate([r for r in results if r['scenario']=='taobao_calibrated']):
        row=[8,9,10,11,12,14,15,16][h]
        if abs(b.get('Итоги','F'+str(row))-x['delta_pp'])>1e-8:raise ValueError('Workbook/Python active results diverge')
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    fd,temporary=tempfile.mkstemp(prefix='ozon_refresh_',suffix='.xlsx',dir=output.parent);os.close(fd)
    try:
        b.save(temporary)
        after=FormulaBook(temporary)
        errors=after.compare_caches()
        if errors:raise ValueError('Refreshed formula caches diverge: '+str(errors[:3]))
        os.replace(temporary,output)
    finally:
        if Path(temporary).exists():Path(temporary).unlink()
    return output
