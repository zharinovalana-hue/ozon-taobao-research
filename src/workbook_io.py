"""Preserving OOXML refresher for this workbook's documented formula subset.

The supplied Excel is the template; this never reconstructs its sheets/styles.
No general Excel compatibility is claimed. Unsupported functions/cycles fail.
Artifact Tool is used for authored layout; this portable path updates inputs and
formula caches for reproducibility outside Codex, including native chart caches.
"""
import math
import operator
import zipfile
from functools import lru_cache
from pathlib import Path
from lxml import etree as ET
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries,get_column_letter

MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
CHART='http://schemas.openxmlformats.org/drawingml/2006/chart'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKGREL='http://schemas.openxmlformats.org/package/2006/relationships'
NS={'s':MAIN,'c':CHART}

def number(x):
    if x is None or x=='':return 0
    if isinstance(x,(int,float)):return x
    raise ValueError('Non-numeric arithmetic operand: '+repr(x))

def flatten(a):
    for v in a:
        if isinstance(v,list):yield from flatten(v)
        else:yield v

@lru_cache(maxsize=20000)
def parse(formula):
    tokens=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE'];pos=0
    def expr(minimum=0):
        nonlocal pos
        if pos>=len(tokens):raise ValueError('Unexpected end of formula')
        t=tokens[pos];pos+=1
        if t.type=='OPERATOR-PREFIX':
            left=('unary',t.value,expr(40))
        elif t.type=='PAREN' and t.subtype=='OPEN':
            left=expr();assert tokens[pos].type=='PAREN' and tokens[pos].subtype=='CLOSE';pos+=1
        elif t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            if not(tokens[pos].type=='FUNC' and tokens[pos].subtype=='CLOSE'):
                while True:
                    args.append(expr())
                    if tokens[pos].type=='SEP':pos+=1;continue
                    break
            if not(tokens[pos].type=='FUNC' and tokens[pos].subtype=='CLOSE'):raise ValueError('Expected function closing parenthesis')
            pos+=1;left=('call',t.value[:-1].upper(),args)
        elif t.type=='OPERAND':
            if t.subtype=='NUMBER':left=('literal',float(t.value))
            elif t.subtype=='TEXT':left=('literal',t.value[1:-1].replace('""','"'))
            elif t.subtype=='LOGICAL':left=('literal',t.value.upper()=='TRUE')
            elif t.subtype=='RANGE':left=('ref',t.value)
            else:raise ValueError('Unsupported operand '+t.value)
        else:raise ValueError('Unsupported formula token '+t.value)
        precedence={'=':10,'<>':10,'<':10,'>':10,'<=':10,'>=':10,'&':15,'+':20,'-':20,'*':30,'/':30,'^':50}
        while pos<len(tokens) and tokens[pos].type=='OPERATOR-INFIX' and precedence.get(tokens[pos].value,-1)>=minimum:
            op=tokens[pos].value;p=precedence[op];pos+=1
            right=expr(p if op=='^' else p+1);left=('binary',op,left,right)
        return left
    tree=expr()
    if pos!=len(tokens):raise ValueError('Unsupported trailing formula tokens: '+formula)
    return tree

class FormulaBook:
    def __init__(self,path):
        self.path=Path(path)
        self.raw=openpyxl.load_workbook(self.path,data_only=False)
        self.cached=openpyxl.load_workbook(self.path,data_only=True)
        self.cells={s.title:{c.coordinate:c.value for row in s for c in row if c.value is not None} for s in self.raw}
        self.memo={};self.stack=set();self.changed={}

    @classmethod
    def from_cells(cls,cells):
        b=cls.__new__(cls);b.path=None;b.cells=cells;b.memo={};b.stack=set();b.changed={};return b

    def set(self,sheet,address,value):
        if sheet not in self.cells:raise ValueError('Unknown worksheet '+sheet)
        address=address.replace('$','')
        self.cells[sheet][address]=value;self.changed[(sheet,address)]=value;self.memo.clear()

    def get(self,sheet,address):
        address=address.replace('$','');key=(sheet,address)
        if key in self.memo:return self.memo[key]
        if key in self.stack:raise ValueError('Circular formula: '+repr(key))
        if sheet not in self.cells:raise ValueError('Unknown worksheet '+sheet)
        v=self.cells[sheet].get(address)
        if isinstance(v,str) and v.startswith('='):
            self.stack.add(key)
            try:v=self._eval(parse(v),sheet)
            except Exception as e:raise ValueError(f'{sheet}!{address}: {e}') from e
            finally:self.stack.remove(key)
        self.memo[key]=v;return v

    def ref(self,text,sheet):
        if '!' in text:
            sheet,text=text.rsplit('!',1)
            if sheet.startswith("'"):sheet=sheet[1:-1].replace("''", "'")
        text=text.replace('$','')
        if ':' not in text:return self.get(sheet,text)
        a,b,c,d=range_boundaries(text)
        return [self.get(sheet,f'{get_column_letter(col)}{row}') for row in range(b,d+1) for col in range(a,c+1)]

    def _eval(self,node,sheet):
        kind=node[0]
        if kind=='literal':return node[1]
        if kind=='ref':return self.ref(node[1],sheet)
        if kind=='unary':return number(self._eval(node[2],sheet))*(-1 if node[1]=='-' else 1)
        if kind=='binary':
            op=node[1];a=self._eval(node[2],sheet);b=self._eval(node[3],sheet)
            if op=='&':return str(a or '')+str(b or '')
            if op in ['=','<>','<','>','<=','>=']:
                if a is None:a='' if isinstance(b,str) else 0
                if b is None:b='' if isinstance(a,str) else 0
                if op in ['=','<>']:
                    same=(a==b) if type(a)==type(b) or isinstance(a,(int,float)) and isinstance(b,(int,float)) else False
                    return same if op=='=' else not same
                return {'<':operator.lt,'>':operator.gt,'<=':operator.le,'>=':operator.ge}[op](a,b)
            a=number(a);b=number(b)
            return {'+':operator.add,'-':operator.sub,'*':operator.mul,'/':operator.truediv,'^':operator.pow}[op](a,b)
        name,args=node[1],node[2]
        if name=='IF':return self._eval(args[1] if self._eval(args[0],sheet) else args[2],sheet)
        if name=='CHOOSE':
            i=self._eval(args[0],sheet)
            if int(i)!=i or not 1<=i<len(args):raise ValueError('CHOOSE index out of range')
            return self._eval(args[int(i)],sheet)
        values=[self._eval(a,sheet) for a in args]
        flat=list(flatten(values));numeric=[v for v in flat if isinstance(v,(int,float)) and not isinstance(v,bool)]
        if name=='AND':return all(bool(v) for v in flat)
        if name=='OR':return any(bool(v) for v in flat)
        if name=='NOT':return not bool(values[0])
        if name=='COUNT':return len(numeric)
        if name=='SUM':return sum(numeric)
        if name=='MIN':return min(numeric) if numeric else 0
        if name=='MAX':return max(numeric) if numeric else 0
        if name=='ABS':return abs(number(values[0]))
        if name=='ISNUMBER':return isinstance(values[0],(int,float)) and not isinstance(values[0],bool)
        if name=='SUMPRODUCT':
            arrays=[v if isinstance(v,list) else [v] for v in values]
            if len({len(a) for a in arrays})!=1:raise ValueError('SUMPRODUCT dimension mismatch')
            return sum(math.prod(number(x) for x in row) for row in zip(*arrays))
        raise ValueError('Unsupported function '+name)

    def calculate(self):
        for sheet,cells in self.cells.items():
            for a,v in cells.items():
                if isinstance(v,str) and v.startswith('='):self.get(sheet,a)
        return self.memo

    def compare_caches(self,tolerance=1e-7):
        errors=[]
        for sheet,cells in self.cells.items():
            for address,formula in cells.items():
                if not isinstance(formula,str) or not formula.startswith('='):continue
                actual=self.get(sheet,address);expected=self.cached[sheet][address].value
                if actual in [None,''] and expected in [None,'']:continue
                if isinstance(actual,(int,float)) and isinstance(expected,(int,float)):
                    if math.isclose(actual,expected,rel_tol=1e-12,abs_tol=tolerance):continue
                elif actual==expected:continue
                errors.append([sheet,address,expected,actual])
        return errors

    @staticmethod
    def _value(cell,value,formula=False):
        for child in list(cell):
            if child.tag in [f'{{{MAIN}}}v',f'{{{MAIN}}}is'] or not formula and child.tag==f'{{{MAIN}}}f':cell.remove(child)
        if value is None:
            cell.attrib.pop('t',None);return
        if isinstance(value,str):
            if formula:
                cell.set('t','str');ET.SubElement(cell,f'{{{MAIN}}}v').text=value
            else:
                cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,f'{{{MAIN}}}is'),f'{{{MAIN}}}t').text=value
        elif isinstance(value,bool):
            cell.set('t','b');ET.SubElement(cell,f'{{{MAIN}}}v').text='1' if value else '0'
        else:
            if not math.isfinite(value):raise ValueError('Cannot save nonfinite result')
            cell.attrib.pop('t',None);ET.SubElement(cell,f'{{{MAIN}}}v').text=format(value,'.17g')

    def save(self,destination):
        """Update existing XML parts only; preserve sheet order/styles/validation/charts."""
        if not self.path:raise ValueError('Need existing workbook as export template')
        self.calculate();destination=Path(destination)
        if destination.resolve()==self.path.resolve():raise ValueError('Use a distinct output path to preserve source template')
        with zipfile.ZipFile(self.path) as zin:
            parts={n:zin.read(n) for n in zin.namelist()}
        wb=ET.fromstring(parts['xl/workbook.xml']);rels=ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        targets={r.get('Id'):r.get('Target') for r in rels}
        paths={s.get('name'):targets[s.get(f'{{{REL}}}id')] for s in wb.find(f'{{{MAIN}}}sheets')}
        for sheet,target in paths.items():
            part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
            xml=ET.fromstring(parts[part]);data=xml.find(f'{{{MAIN}}}sheetData')
            mapping={c.get('r'):c for row in data for c in row if c.tag==f'{{{MAIN}}}c'}
            rowmap={int(r.get('r')):r for r in data}
            for address,value in self.cells[sheet].items():
                formula=isinstance(value,str) and value.startswith('=')
                if not formula and (sheet,address) not in self.changed:continue
                cell=mapping.get(address)
                if cell is None:
                    _,r,_,_=range_boundaries(address)
                    row=rowmap.get(r)
                    if row is None:row=ET.SubElement(data,f'{{{MAIN}}}row',r=str(r));rowmap[r]=row
                    cell=ET.SubElement(row,f'{{{MAIN}}}c',r=address);mapping[address]=cell
                if formula:
                    old=cell.find(f'{{{MAIN}}}f')
                    if old is None:old=ET.SubElement(cell,f'{{{MAIN}}}f')
                    old.text=value[1:]
                    self._value(cell,self.get(sheet,address),True)
                else:self._value(cell,value)
            for row in data:
                row[:]=sorted(row,key=lambda c:range_boundaries(c.get('r'))[0])
            data[:]=sorted(data,key=lambda r:int(r.get('r')))
            parts[part]=ET.tostring(xml,encoding='utf-8',xml_declaration=True,standalone=True)
        calc=wb.find(f'{{{MAIN}}}calcPr')
        if calc is None:calc=ET.SubElement(wb,f'{{{MAIN}}}calcPr')
        calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
        parts['xl/workbook.xml']=ET.tostring(wb,encoding='utf-8',xml_declaration=True,standalone=True)
        for name,blob in list(parts.items()):
            if not name.startswith('xl/charts/chart') or not name.endswith('.xml'):continue
            chart=ET.fromstring(blob)
            for tag,cacheTag in [('numRef','numCache'),('strRef','strCache')]:
                for ref in chart.findall('.//c:'+tag,NS):
                    f=ref.find('c:f',NS)
                    if f is None:continue
                    values=self.ref(f.text,'Итоги');values=values if isinstance(values,list) else [values]
                    old=ref.find('c:'+cacheTag,NS)
                    if old is not None:ref.remove(old)
                    cache=ET.SubElement(ref,f'{{{CHART}}}{cacheTag}')
                    if tag=='numRef':ET.SubElement(cache,f'{{{CHART}}}formatCode').text='General'
                    ET.SubElement(cache,f'{{{CHART}}}ptCount',val=str(len(values)))
                    for i,val in enumerate(values):
                        point=ET.SubElement(cache,f'{{{CHART}}}pt',idx=str(i))
                        ET.SubElement(point,f'{{{CHART}}}v').text=str(val if val is not None else '')
            parts[name]=ET.tostring(chart,encoding='utf-8',xml_declaration=True,standalone=True)
        destination.parent.mkdir(parents=True,exist_ok=True)
        with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_DEFLATED) as zout:
            for name,blob in parts.items():zout.writestr(name,blob)
